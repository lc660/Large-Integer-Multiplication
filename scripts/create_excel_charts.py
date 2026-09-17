#!/usr/bin/env python3
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
AWS_RESULTS = ROOT / "results" / "aws_ec2_multiply_benchmark.csv"
ECE_RESULTS = ROOT / "results" / "ece_cluster_multiply_benchmark.csv"
OUTPUT = ROOT / "results" / "large_integer_multiplication_charts.xlsx"


def load_machine_results(path):
    rows = {}
    with path.open(newline="") as f:
        for row in csv.DictReader(f):
            digits = int(row["digits"])
            rows.setdefault(digits, {})[row["method"]] = float(row["flops_per_second"])
    return rows


def style_header(ws, row=1):
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


def add_chart_sheet(wb, data_ws, title, method_column_start, sheet_name):
    ws = wb.create_sheet(sheet_name)
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True)
    ws["A3"] = "Clustered columns compare ECE Cluster and AWS EC2 for each input length."
    ws["A3"].font = Font(italic=True, color="666666")

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = title
    chart.y_axis.title = "flops/s"
    chart.x_axis.title = "Input length, n digits"
    chart.width = 24
    chart.height = 13
    chart.grouping = "clustered"
    chart.overlap = 0
    chart.legend.position = "r"

    max_row = data_ws.max_row
    data = Reference(
        data_ws,
        min_col=method_column_start,
        max_col=method_column_start + 1,
        min_row=1,
        max_row=max_row,
    )
    cats = Reference(data_ws, min_col=1, min_row=2, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, "A5")

    return ws


def main():
    aws_results = load_machine_results(AWS_RESULTS)
    ece_results = load_machine_results(ECE_RESULTS)
    all_digits = sorted(set(aws_results) | set(ece_results))

    wb = Workbook()
    data_ws = wb.active
    data_ws.title = "Data"
    data_ws.append([
        "digits",
        "ECE Cluster O(n^2) flops/s",
        "AWS EC2 O(n^2) flops/s",
        "ECE Cluster O(n log2 n) flops/s",
        "AWS EC2 O(n log2 n) flops/s",
    ])

    for digits in all_digits:
        data_ws.append([
            digits,
            ece_results.get(digits, {}).get("quadratic"),
            aws_results.get(digits, {}).get("quadratic"),
            ece_results.get(digits, {}).get("fft"),
            aws_results.get(digits, {}).get("fft"),
        ])

    style_header(data_ws)
    data_ws.freeze_panes = "A2"
    for col in range(1, 6):
        data_ws.column_dimensions[get_column_letter(col)].width = 28
    for row in data_ws.iter_rows(min_row=2, min_col=2, max_col=5):
        for cell in row:
            cell.number_format = "0.00E+00"

    data_ws["G1"] = "How to use"
    data_ws["G1"].fill = PatternFill("solid", fgColor="D9EAF7")
    data_ws["G1"].font = Font(bold=True)
    data_ws["G2"] = "Source files: aws_ec2_multiply_benchmark.csv and ece_cluster_multiply_benchmark.csv."
    data_ws["G3"] = "Each chart contains ECE Cluster and AWS EC2 clustered columns."
    data_ws.column_dimensions["G"].width = 56

    add_chart_sheet(wb, data_ws, "O(n^2) Multiplication: flops/s vs Problem Size", 2, "O(n^2) Chart")
    add_chart_sheet(wb, data_ws, "O(n log2 n) FFT Multiplication: flops/s vs Problem Size", 4, "O(n log n) Chart")

    wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
